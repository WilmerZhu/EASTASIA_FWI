#!/usr/bin/env python3
"""
Markdown 主稿 → 《中国科学：地球科学》投稿格式 Word

用法:
    python md2docx.py                    # 默认转换本目录 md
    python md2docx.py --md xxx.md --out yyy.docx

Markdown 约定:
    ---            YAML 头（title/authors/affiliations/corresponding/figdir）
    **摘要** ...    加粗行首标签段（摘要/关键词/项目资助/致谢）
    # 1 引言       一级标题（黑体小四，不加粗）
    ## 3.1 ...     二级标题（宋体加粗五号）
    $$...\\tag{n}$$ 单列公式，右端编号 (n)
    ![](fig1.jpg)  插图，路径相对 figdir
    **图1** ...     图题（居中加粗）
    **表1** ...     表题（居中加粗）
    | a | b |      三线表
    注：...         表注（小五号）
    <!-- ... -->   注释块，整段跳过

版式依据投稿模板: 正文宋体/Times New Roman 五号、1.5 倍行距、首行缩进 2 字；
题目黑体三号不加粗，一级标题黑体小四不加粗，二级、三级标题宋体加粗五号；
摘要/关键词/项目资助/致谢的行首标签黑体不加粗，整段左对齐且不缩进；
通讯作者行左对齐，作者与单位行居中。
"""
from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph
from lxml import etree

HERE = Path(__file__).resolve().parent

FONT_EN = "Times New Roman"
FONT_SONG = "宋体"
FONT_HEI = "黑体"

SZ_TITLE = Pt(16)      # 三号
SZ_H1 = Pt(12)         # 小四
SZ_BODY = Pt(10.5)     # 五号
SZ_NOTE = Pt(9)        # 小五
LINE_15 = 1.5
INDENT = Pt(21)        # 首行缩进 2 个五号汉字

FIG_WIDTH_CM = 14.5
TAB_RIGHT_TWIPS = "9026"   # A4 去 2.54 cm 页边距后的正文右缘

M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NSMAP = {"m": M_NS, "w": W_NS}


# ==================== 基础排版工具 ====================

def set_run_font(run, *, east_asia: str = FONT_SONG, ascii_font: str = FONT_EN,
                 size=SZ_BODY, bold: bool = False, italic: bool = False) -> None:
    run.bold = bold
    run.italic = italic
    run.font.size = size
    run.font.name = ascii_font
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr, val in (("w:ascii", ascii_font), ("w:hAnsi", ascii_font),
                      ("w:eastAsia", east_asia), ("w:cs", ascii_font)):
        rFonts.set(qn(attr), val)


def set_paragraph_format(p, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_indent=None,
                         space_before: float = 0, space_after: float = 0,
                         line=LINE_15, keep_with_next: bool = False,
                         snap_to_grid: bool = True) -> None:
    pf = p.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    pf.first_line_indent = first_indent
    pf.widow_control = True
    pf.keep_with_next = keep_with_next
    pPr = p._p.get_or_add_pPr()
    snap = pPr.find(qn("w:snapToGrid"))
    if snap is None:
        snap = OxmlElement("w:snapToGrid")
        pPr.append(snap)
    snap.set(qn("w:val"), "true" if snap_to_grid else "false")


# 行内 **加粗** / *斜体* 解析
_INLINE = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*))")


def add_rich_text(p, text: str, *, east_asia: str = FONT_SONG, size=SZ_BODY,
                  base_bold: bool = False) -> None:
    """写入一段文本，处理 **粗体**、*斜体* 与转义符。"""
    for chunk in _INLINE.split(text):
        if not chunk:
            continue
        bold, italic = base_bold, False
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            chunk, bold = chunk[2:-2], True
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            chunk, italic = chunk[1:-1], True
        chunk = chunk.replace("\\*", "*").replace("\\|", "|").replace("\\_", "_")
        run = p.add_run(chunk)
        set_run_font(run, east_asia=east_asia, size=size, bold=bold, italic=italic)


def force_auto_line_spacing(p, before: int = 120, after: int = 40) -> None:
    """插图段必须用自动行距，否则文档网格会把图片裁成一行高。"""
    pPr = p._p.get_or_add_pPr()
    spacing = pPr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        pPr.append(spacing)
    spacing.set(qn("w:before"), str(before))
    spacing.set(qn("w:after"), str(after))
    spacing.set(qn("w:line"), "240")
    spacing.set(qn("w:lineRule"), "auto")
    old = pPr.find(qn("w:ind"))
    if old is not None:
        pPr.remove(old)
    ind = OxmlElement("w:ind")
    for attr in ("w:firstLine", "w:left", "w:right"):
        ind.set(qn(attr), "0")
    pPr.append(ind)


def set_cell_border(cell, edges: Dict[str, Dict[str, str]]) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge, spec in edges.items():
        old = borders.find(qn(f"w:{edge}"))
        if old is not None:
            borders.remove(old)
        el = OxmlElement(f"w:{edge}")
        for k, v in spec.items():
            el.set(qn(f"w:{k}"), v)
        borders.append(el)


NO_BORDER = {"val": "nil", "sz": "0", "space": "0", "color": "auto"}
LINE_THICK = {"val": "single", "sz": "12", "space": "0", "color": "000000"}
LINE_THIN = {"val": "single", "sz": "6", "space": "0", "color": "000000"}


# ==================== 公式：LaTeX → OMML ====================

def latex_to_omath(latex_list: List[str]) -> List[etree._Element]:
    """用 pandoc 批量把 LaTeX 转成可在 Word 中编辑的 m:oMath 节点。"""
    if not latex_list:
        return []
    md = "\n\n".join(f"$${e}$$" for e in latex_list)
    with tempfile.TemporaryDirectory() as td:
        src, out = Path(td) / "eq.md", Path(td) / "eq.docx"
        src.write_text(md, encoding="utf-8")
        subprocess.check_call(
            ["pandoc", str(src), "-f", "markdown", "-t", "docx", "-o", str(out)]
        )
        with zipfile.ZipFile(out) as zf:
            root = etree.fromstring(zf.read("word/document.xml"))
    maths = root.xpath("//m:oMath", namespaces=NSMAP)
    if len(maths) != len(latex_list):
        raise RuntimeError(f"公式转换数量不符: {len(maths)} != {len(latex_list)}")
    cleaned = []
    for m in maths:
        node = deepcopy(m)
        etree.cleanup_namespaces(node, top_nsmap={"m": M_NS, "w": W_NS})
        cleaned.append(node)
    return cleaned


# ==================== Markdown 解析 ====================

BLOCK_COMMENT = re.compile(r"<!--.*?-->", re.S)
RE_FIG = re.compile(r"^!\[.*?\]\((.+?)\)\s*$")
RE_EQ = re.compile(r"^\$\$(.+?)\$\$\s*$", re.S)
RE_TAG = re.compile(r"\\tag\{([^}]+)\}")
RE_H = re.compile(r"^(#{1,4})\s+(.*)$")
RE_TABLE_SEP = re.compile(r"^\|[\s:\-\|]+\|$")
RE_LABEL = re.compile(r"^\*\*(摘要|关键词|项目资助|致谢)\*\*\s*(.*)$")
RE_CAPTION = re.compile(r"^\*\*(图|表)\s*(S?\d+)\*\*\s*(.*)$")   # 支持正文「表1」与补充材料「表S1」


def parse_front_matter(text: str) -> Tuple[Dict[str, Any], str]:
    if not text.startswith("---"):
        return {}, text
    end = text.find("\n---", 3)
    if end < 0:
        return {}, text
    head, body = text[3:end], text[end + 4:]
    meta: Dict[str, Any] = {}
    key = None
    for line in head.splitlines():
        if not line.strip():
            continue
        if line.lstrip().startswith("- ") and key:
            meta.setdefault(key, [])
            if isinstance(meta[key], list):
                meta[key].append(line.lstrip()[2:].strip())
            continue
        if ":" in line:
            key, _, val = line.partition(":")
            key = key.strip()
            val = val.strip().strip('"').strip("'")
            meta[key] = val if val else []
    return meta, body


def split_table(lines: List[str], i: int) -> Tuple[Optional[List[List[str]]], int]:
    """从 lines[i] 起解析 markdown 管道表，返回 (rows, 下一行索引)。"""
    if not lines[i].lstrip().startswith("|"):
        return None, i
    if i + 1 >= len(lines) or not RE_TABLE_SEP.match(lines[i + 1].strip()):
        return None, i
    rows: List[List[str]] = []
    j = i
    while j < len(lines) and lines[j].lstrip().startswith("|"):
        if RE_TABLE_SEP.match(lines[j].strip()):
            j += 1
            continue
        raw = lines[j].strip().strip("|")
        cells = [c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", raw)]
        rows.append(cells)
        j += 1
    return rows, j


# ==================== 写入 docx ====================

class DocBuilder:
    def __init__(self, meta: Dict[str, Any], figdir: Path):
        self.meta = meta
        self.figdir = figdir
        self.doc = Document()
        self._setup_page()
        self.missing_figs: List[str] = []

    def _setup_page(self) -> None:
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        for attr in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, attr, Cm(2.54))
        style = self.doc.styles["Normal"]
        style.font.name = FONT_EN
        style.font.size = SZ_BODY
        style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_SONG)

    # ---- 段落类型 ----

    def para(self, **kw) -> Paragraph:
        p = self.doc.add_paragraph()
        set_paragraph_format(p, **kw)
        return p

    def title_block(self) -> None:
        p = self.para(align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=6)
        add_rich_text(p, self.meta.get("title", ""), east_asia=FONT_HEI,
                      size=SZ_TITLE)

        authors = self.meta.get("authors", "")
        if authors:
            p = self.para(align=WD_ALIGN_PARAGRAPH.CENTER, space_after=3)
            self._add_with_superscript(p, authors)

        for aff in self.meta.get("affiliations", []) or []:
            p = self.para(align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
            add_rich_text(p, aff)

        corr = self.meta.get("corresponding", "")
        if corr:
            p = self.para(align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=None,
                          space_before=3, space_after=8)
            add_rich_text(p, corr)

    @staticmethod
    def _add_with_superscript(p, text: str) -> None:
        """处理作者行的 ^1^ 上标标记。"""
        for part in re.split(r"(\^[^\^]+\^)", text):
            if not part:
                continue
            sup = part.startswith("^") and part.endswith("^") and len(part) > 2
            run = p.add_run(part[1:-1] if sup else part)
            set_run_font(run, size=SZ_BODY)
            run.font.superscript = sup

    def labeled(self, label: str, rest: str) -> None:
        """摘要/关键词/项目资助/致谢：黑体不加粗的行首标签，整段左对齐、不缩进。"""
        p = self.para(align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=None,
                      space_before=3, space_after=3)
        run = p.add_run(label + "  ")
        set_run_font(run, east_asia=FONT_HEI, size=SZ_BODY)
        add_rich_text(p, rest)

    def heading(self, level: int, text: str) -> None:
        if level == 1:
            p = self.para(align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=None,
                          space_before=12, space_after=6, keep_with_next=True)
            add_rich_text(p, text, east_asia=FONT_HEI, size=SZ_H1)
        else:
            p = self.para(align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=None,
                          space_before=8, space_after=3, keep_with_next=True)
            add_rich_text(p, text, east_asia=FONT_SONG, size=SZ_BODY, base_bold=True)

    def body(self, text: str) -> None:
        p = self.para(first_indent=INDENT)
        add_rich_text(p, text)

    def note(self, text: str) -> None:
        p = self.para(first_indent=None, space_before=2, space_after=8)
        add_rich_text(p, text, size=SZ_NOTE)

    def reference(self, text: str) -> None:
        """参考文献：悬挂缩进，便于阅读；Zotero 后续可整体替换。"""
        p = self.para(first_indent=Pt(-21), space_after=2)
        p.paragraph_format.left_indent = Pt(21)
        add_rich_text(p, text)

    def caption(self, text: str) -> None:
        p = self.para(align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=None,
                      space_before=3, space_after=8)
        add_rich_text(p, text, base_bold=True)

    def figure(self, rel_path: str) -> None:
        path = self.figdir / rel_path
        p = self.para(align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=Pt(0),
                      space_before=6, space_after=2, line=1.0,
                      keep_with_next=True, snap_to_grid=False)
        force_auto_line_spacing(p)
        if not path.exists():
            self.missing_figs.append(str(path))
            add_rich_text(p, f"【缺图：{rel_path}】")
            return
        p.add_run().add_picture(str(path), width=Cm(FIG_WIDTH_CM))

    def equation(self, omml: etree._Element, number: Optional[str]) -> None:
        p = self.para(align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=None,
                      space_before=6, space_after=6, line=1.15, snap_to_grid=False)
        pPr = p._p.get_or_add_pPr()
        old = pPr.find(qn("w:tabs"))
        if old is not None:
            pPr.remove(old)
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "right")
        tab.set(qn("w:pos"), TAB_RIGHT_TWIPS)
        tabs.append(tab)
        pPr.append(tabs)
        p._p.append(deepcopy(omml))
        if number:
            p.add_run()._r.append(OxmlElement("w:tab"))
            run = p.add_run(f"({number})")
            set_run_font(run, size=SZ_BODY)

    def table(self, rows: List[List[str]]) -> None:
        """三线表：顶线与底线粗、表头下细线，其余无框。"""
        n_col = max(len(r) for r in rows)
        table = self.doc.add_table(rows=len(rows), cols=n_col)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        last = len(rows) - 1
        for i, row in enumerate(rows):
            for j in range(n_col):
                cell = table.cell(i, j)
                cell.paragraphs[0].text = ""
                p = cell.paragraphs[0]
                set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.CENTER,
                                     first_indent=None, line=1.0, snap_to_grid=False)
                add_rich_text(p, row[j] if j < len(row) else "",
                              size=SZ_NOTE, base_bold=(i == 0))
                edges = {"left": NO_BORDER, "right": NO_BORDER,
                         "top": NO_BORDER, "bottom": NO_BORDER}
                if i == 0:
                    edges["top"] = LINE_THICK
                    edges["bottom"] = LINE_THIN
                if i == last:
                    edges["bottom"] = LINE_THICK
                set_cell_border(cell, edges)


# ==================== 主流程 ====================

def build(md_path: Path, out_path: Path) -> None:
    raw = md_path.read_text(encoding="utf-8")
    raw = BLOCK_COMMENT.sub("", raw)
    meta, body = parse_front_matter(raw)

    figdir = (md_path.parent / meta.get("figdir", "../figures")).resolve()
    builder = DocBuilder(meta, figdir)
    builder.title_block()

    lines = body.splitlines()
    # 先收集全部公式，一次性调用 pandoc
    latex_specs: List[Tuple[str, Optional[str]]] = []
    for line in lines:
        m = RE_EQ.match(line.strip())
        if m:
            expr = m.group(1).strip()
            tag = RE_TAG.search(expr)
            latex_specs.append((RE_TAG.sub("", expr).strip(),
                                tag.group(1) if tag else None))
    ommls = latex_to_omath([e for e, _ in latex_specs])
    eq_iter = iter(zip(ommls, [n for _, n in latex_specs]))

    in_refs = False
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        rows, nxt = split_table(lines, i)
        if rows:
            builder.table(rows)
            i = nxt
            continue

        m = RE_H.match(stripped)
        if m:
            level, text = len(m.group(1)), m.group(2).strip()
            in_refs = text.startswith("参考文献")
            builder.heading(level, text)
            i += 1
            continue

        m = RE_EQ.match(stripped)
        if m:
            omml, number = next(eq_iter)
            builder.equation(omml, number)
            i += 1
            continue

        m = RE_FIG.match(stripped)
        if m:
            builder.figure(m.group(1).strip())
            i += 1
            continue

        m = RE_LABEL.match(stripped)
        if m:
            builder.labeled(m.group(1), m.group(2))
            i += 1
            continue

        m = RE_CAPTION.match(stripped)
        if m:
            builder.caption(f"{m.group(1)}{m.group(2)}  {m.group(3)}")
            i += 1
            continue

        if stripped.startswith("注："):
            builder.note(stripped)
            i += 1
            continue

        if in_refs:
            builder.reference(stripped)
        else:
            builder.body(stripped)
        i += 1

    if out_path.exists():
        backup_dir = out_path.parent / "backup"
        backup_dir.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        bak = backup_dir / f"{out_path.stem}_{stamp}{out_path.suffix}"
        shutil.copy2(out_path, bak)
        print(f"已备份旧稿: {bak.name}")

    builder.doc.save(str(out_path))
    print(f"✅ 已生成: {out_path}")
    if builder.missing_figs:
        print("⚠️ 缺失插图:")
        for f in builder.missing_figs:
            print(f"   {f}")


def main() -> None:
    ap = argparse.ArgumentParser(description="Markdown 主稿转投稿格式 Word")
    ap.add_argument("--md", default=str(HERE / "东亚地震学模型对比与检验.md"))
    ap.add_argument("--out", default=str(HERE / "东亚地震学模型对比与检验.docx"))
    args = ap.parse_args()

    md_path, out_path = Path(args.md), Path(args.out)
    if not md_path.exists():
        sys.exit(f"❌ 找不到 Markdown 源文件: {md_path}")
    if shutil.which("pandoc") is None:
        sys.exit("❌ 需要 pandoc 来生成 Word 公式对象，请先安装")
    build(md_path, out_path)


if __name__ == "__main__":
    main()
